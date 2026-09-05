"""Safe source ingestion for staff-managed after-sales knowledge."""

from __future__ import annotations

import io
import ipaddress
import re
import socket
from html import unescape
from urllib.parse import urljoin, urlparse

import requests
from django.conf import settings


MAX_SOURCE_BYTES = 12 * 1024 * 1024
ALLOWED_FILE_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}


class KnowledgeIngestError(ValueError):
    """A safe, user-facing source parsing error."""


def _clean_text(value: str) -> str:
    value = unescape(value).replace("\x00", "")
    value = re.sub(r"\r\n?", "\n", value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _read_limited(uploaded_file) -> bytes:
    size = getattr(uploaded_file, "size", None)
    if size is not None and size > MAX_SOURCE_BYTES:
        raise KnowledgeIngestError("文件不能超过 12 MB。")
    chunks: list[bytes] = []
    total = 0
    for chunk in uploaded_file.chunks():
        total += len(chunk)
        if total > MAX_SOURCE_BYTES:
            raise KnowledgeIngestError("文件不能超过 12 MB。")
        chunks.append(chunk)
    return b"".join(chunks)


def extract_uploaded_file(uploaded_file) -> tuple[str, str]:
    """Extract text from txt, markdown, Word or PDF without executing content."""

    import os

    name = str(getattr(uploaded_file, "name", ""))
    extension = os.path.splitext(name.lower())[1]
    if extension not in ALLOWED_FILE_EXTENSIONS:
        raise KnowledgeIngestError("仅支持 .txt、.md、.docx 和 .pdf 文件。")
    raw = _read_limited(uploaded_file)

    if extension in {".txt", ".md"}:
        text = raw.decode("utf-8-sig", errors="replace")
    elif extension == ".docx":
        try:
            from docx import Document

            document = Document(io.BytesIO(raw))
        except Exception as exc:  # pragma: no cover - parser library owns details
            raise KnowledgeIngestError("Word 文档解析失败，请确认文件没有损坏。") from exc
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(paragraphs)
    else:
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:  # pragma: no cover - parser library owns details
            raise KnowledgeIngestError("PDF 文档解析失败，请确认文件没有损坏或是扫描图片。") from exc

    text = _clean_text(text)
    if len(text) < 2:
        raise KnowledgeIngestError("文档没有提取到可用文字。")
    return text, extension.lstrip(".")


def _validate_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise KnowledgeIngestError("网页地址必须以 http:// 或 https:// 开头。")
    if parsed.username or parsed.password:
        raise KnowledgeIngestError("网页地址不能包含登录凭据。")
    try:
        addresses = {
            ipaddress.ip_address(item[4][0])
            for item in socket.getaddrinfo(parsed.hostname, None, type=socket.SOCK_STREAM)
        }
    except (OSError, ValueError) as exc:
        raise KnowledgeIngestError("网页域名无法解析。") from exc
    if any(
        address.is_private
        or address.is_loopback
        or address.is_link_local
        or address.is_reserved
        or address.is_unspecified
        for address in addresses
    ):
        raise KnowledgeIngestError("为保护系统安全，不允许读取内网或本机地址。")


def extract_webpage(url: str) -> tuple[str, str]:
    """Fetch a public HTML page and retain readable text only."""

    current_url = url.strip()
    response = None
    for _ in range(4):
        _validate_public_url(current_url)
        try:
            response = requests.get(
                current_url,
                timeout=settings.KNOWLEDGE_WEB_TIMEOUT_SECONDS,
                headers={"User-Agent": "AfterSalesKnowledgeBot/1.0"},
                allow_redirects=False,
                stream=True,
            )
        except requests.RequestException as exc:
            raise KnowledgeIngestError("网页读取失败，请检查地址是否可访问。") from exc
        if response.is_redirect or response.is_permanent_redirect:
            location = response.headers.get("Location")
            if not location:
                response.close()
                raise KnowledgeIngestError("网页重定向地址无效。")
            current_url = urljoin(current_url, location)
            response.close()
            continue
        break
    else:
        raise KnowledgeIngestError("网页重定向次数过多。")
    if response is None:
        raise KnowledgeIngestError("网页读取失败，请检查地址是否可访问。")
    try:
        response.raise_for_status()
    except requests.RequestException as exc:
        response.close()
        raise KnowledgeIngestError("网页读取失败，请检查地址是否可访问。") from exc
    content_length = response.headers.get("Content-Length")
    try:
        if content_length and int(content_length) > MAX_SOURCE_BYTES:
            response.close()
            raise KnowledgeIngestError("网页内容不能超过 12 MB。")
    except ValueError:
        pass
    content_type = response.headers.get("Content-Type", "").lower()
    if content_type and "html" not in content_type and "text/" not in content_type:
        response.close()
        raise KnowledgeIngestError("网页地址没有返回 HTML 或文本内容。")

    try:
        raw_content = bytearray()
        for chunk in response.iter_content(chunk_size=64 * 1024):
            if not chunk:
                continue
            raw_content.extend(chunk)
            if len(raw_content) > MAX_SOURCE_BYTES:
                raise KnowledgeIngestError("网页内容不能超过 12 MB。")
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(bytes(raw_content), "html.parser")
        for node in soup(["script", "style", "noscript", "svg", "nav", "footer"]):
            node.decompose()
        title = _clean_text(soup.title.get_text(" ") if soup.title else "")
        text = _clean_text(soup.get_text("\n"))
    except Exception as exc:  # pragma: no cover - parser library owns details
        if isinstance(exc, KnowledgeIngestError):
            raise
        raise KnowledgeIngestError("网页正文解析失败。") from exc
    finally:
        response.close()
    if len(text) < 2:
        raise KnowledgeIngestError("网页没有提取到可用文字。")
    return text, title


def source_type_for(*, uploaded_file=None, source_url: str = "", content: str = "") -> str:
    if uploaded_file is not None:
        return "FILE"
    if source_url:
        return "WEBPAGE"
    if content.strip():
        return "TEXT"
    raise KnowledgeIngestError("请上传文件、填写网页地址或输入知识内容。")
